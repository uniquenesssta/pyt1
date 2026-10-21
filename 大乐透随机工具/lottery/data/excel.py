"""Excel 读取适配层，共用文本号码校验。"""
from dataclasses import replace
from pathlib import Path
from ..core.models import check_cancel
from .importing import Dataset, ImportIssue, parse_line


def workbook_rows(path):
    if Path(path).suffix.lower() == '.xls':
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError('缺少 xlrd，请通过启动.cmd 安装 Excel 依赖') from exc
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for sheet in book.sheets():
                for row in range(sheet.nrows):
                    values = []
                    for cell in sheet.row(row):
                        if cell.ctype in (xlrd.XL_CELL_DATE, xlrd.XL_CELL_BOOLEAN, xlrd.XL_CELL_ERROR):
                            values.append('[日期/布尔/错误单元格]')
                        else:
                            values.append(cell.value)
                    yield sheet.name, row + 1, values
        finally:
            book.release_resources()
    else:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError('缺少 openpyxl，请通过启动.cmd 安装 Excel 依赖') from exc
        with open(path, 'rb') as stream:
            book = load_workbook(stream, read_only=True, data_only=False)
            try:
                for sheet in book.worksheets:
                    for number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                        yield sheet.title, number, row
            finally:
                book.close()


def load_excel(path, cancel=None, progress=None):
    result = Dataset(source=str(path), encoding='Excel')
    seen = set()
    rows = workbook_rows(path)
    try:
        for sheet, line, values in rows:
            check_cancel(cancel)
            if progress and line % 256 == 0:
                progress(f'正在读取 {sheet}：第 {line:,} 行')
            cells = ['' if v is None else str(int(v)) if type(v) is float and v.is_integer() else str(v).strip() for v in values]
            while cells and not cells[-1]:
                cells.pop()
            if not cells or not any(cells):
                result.ignored += 1
                continue
            text = ' | '.join(cells)
            try:
                if any(not c for c in cells):
                    raise ValueError('行内有空单元格，请补全号码或移除空列')
                draw = replace(parse_line(text, line), sheet=sheet)
                key = (draw.issue, draw.ticket)
                if key in seen:
                    result.duplicates += 1
                seen.add(key)
                result.draws.append(draw)
            except ValueError as exc:
                result.issues.append(ImportIssue(line, f'工作表「{sheet}」：{exc}', text))
    finally:
        rows.close()
    check_cancel(cancel)
    return result
